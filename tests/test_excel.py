import io
import unittest

import pandas as pd
from openpyxl import load_workbook

import config
from services import excel_service as xs
from tests.helpers import make_services


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.s = make_services(sample=True)
        self.addCleanup(self.s.cleanup)

    def _xlsx(self, df):
        b = io.BytesIO()
        df.to_excel(b, index=False)
        b.seek(0)
        b.name = "in.xlsx"
        return b

    def test_template_roundtrip(self):
        tpl = io.BytesIO(xs.build_template())
        tpl.name = "t.xlsx"
        df, _ = xs.read_books_file(tpl)
        self.assertEqual(list(df.columns), config.BOOK_COLUMNS)
        self.assertEqual(len(df), 1)

    def test_export_has_four_sheets_and_roundtrips(self):
        data = xs.export_workbook(self.s.cat.list_books(), self.s.users.get_history(), None, self.s.cat.list_scan_logs())
        self.assertEqual(load_workbook(io.BytesIO(data)).sheetnames,
                         ["Book Catalogue", "Reading History", "Recommendations", "Scan Logs"])
        back, _ = xs.read_books_file(io.BytesIO(data), "x.xlsx")
        self.assertEqual(len(back), 52)

    def test_export_empty_data_still_valid(self):
        self.assertEqual(len(load_workbook(io.BytesIO(xs.export_workbook(None, None, None, None))).sheetnames), 4)

    def test_aliases_missing_values_duplicates_bad_isbn(self):
        df = pd.DataFrame({"Book Title": ["A Book", "a book", "", "Another"], "Writer": ["X Y", "X Y", "Z", None],
                           "ISBN": ["978-0-306-40615-8", "", "", "978-0-306-40615-7"], "Junk": [1, 2, 3, 4]})
        out, issues = xs.read_books_file(self._xlsx(df))
        self.assertEqual(list(out["title"]), ["A Book", "Another"])
        self.assertEqual(out.loc[0, "isbn"], "")                 # invalid checksum dropped
        self.assertEqual(out.loc[1, "isbn"], "9780306406157")    # valid ISBN normalised
        text = " ".join(issues)
        for needle in ("no title", "duplicate", "invalid ISBN", "Junk"):
            self.assertIn(needle, text)

    def test_missing_title_column(self):
        with self.assertRaises(xs.ExcelFormatError):
            xs.read_books_file(self._xlsx(pd.DataFrame({"author": ["x"]})))

    def test_invalid_and_unsupported_files(self):
        bad = io.BytesIO(b"this is not a zip")
        bad.name = "broken.xlsx"
        with self.assertRaises(xs.ExcelFormatError):
            xs.read_books_file(bad)
        txt = io.BytesIO(b"x")
        txt.name = "notes.pdf"
        with self.assertRaises(xs.ExcelFormatError):
            xs.read_books_file(txt)

    def test_csv_and_import_into_catalogue_is_non_destructive(self):
        csv = io.BytesIO(b"title,author,genre\nDune,Frank Herbert,Fantasy\nNew Book,Someone,Mystery\n")
        csv.name = "b.csv"
        df, _ = xs.read_books_file(csv)
        did = self.s.cat.find_duplicate("Dune", "Frank Herbert")
        before = self.s.cat.get_book(did)["genre"]
        counts = self.s.cat.add_many(df.to_dict("records"))
        self.assertEqual((counts["added"], counts["duplicate"] + counts["merged"]), (1, 1))
        self.assertEqual(self.s.cat.get_book(did)["genre"], before)

    def test_history_import(self):
        h = pd.DataFrame({"user": ["U1"], "title": ["Dune"], "status": ["READ"], "stars": [5]})
        out, _ = xs.read_history_file(self._xlsx(h))
        self.assertEqual((out.loc[0, "user_id"], out.loc[0, "status"], out.loc[0, "rating"]), ("U1", "read", "5"))
        with self.assertRaises(xs.ExcelFormatError):
            xs.read_history_file(self._xlsx(pd.DataFrame({"foo": [1]})))


if __name__ == "__main__":
    unittest.main()
