"""Convert docs/PROJECT_REPORT.md to a formal Word report.
Formatting: Times New Roman; body 12 pt; Heading 1 = 28 pt, Heading 2 = 18 pt, Heading 3 = 16 pt (constant).
Usage: pip install python-docx ; python docs/build_report_docx.py"""
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).parent
FONT = "Times New Roman"
SIZES = {"Heading 1": 28, "Heading 2": 18, "Heading 3": 16}


def set_font(run, size, bold=None, mono=False):
    name = "Courier New" if mono else FONT
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if bold is not None:
        run.bold = bold


def add_inline(par, text, size=12):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_font(par.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("`"):
            set_font(par.add_run(part[1:-1]), 11, mono=True)
        elif part.startswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size); r.italic = True
        else:
            set_font(par.add_run(part), size)


def main():
    doc = Document()
    for name, size in {"Normal": 12, **SIZES}.items():
        st = doc.styles[name]
        st.font.name, st.font.size = FONT, Pt(size)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        if name != "Normal":
            st.font.bold = True
            st.font.color.rgb = None
    lines = (HERE / "PROJECT_REPORT.md").read_text(encoding="utf-8").splitlines()
    i, in_code, code, table = 0, False, [], []

    def flush_table():
        nonlocal table
        rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table if not re.match(r"^\|[-| :]+\|$", r.strip())]
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row[: len(rows[0])]):
                    p = t.cell(ri, ci).paragraphs[0]
                    add_inline(p, cell if ri else f"**{cell}**", 11)
        table = []

    for line in lines:
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                set_font(p.add_run("\n".join(code)), 10, mono=True)
                code = []
            in_code = not in_code
            continue
        if in_code:
            code.append(line)
            continue
        if line.startswith("|"):
            table.append(line)
            continue
        if table:
            flush_table()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            h = doc.add_heading(level=len(m.group(1)))
            r = h.add_run(re.sub(r"[*`]", "", m.group(2)))
            set_font(r, SIZES[f"Heading {len(m.group(1))}"], bold=True)
        elif re.match(r"^(\d+\.|-) ", line):
            p = doc.add_paragraph(style="List Number" if line[0].isdigit() else "List Bullet")
            add_inline(p, re.sub(r"^(\d+\.|-) ", "", line))
        else:
            add_inline(doc.add_paragraph(), line)
    if table:
        flush_table()
    out = HERE / "PROJECT_REPORT.docx"
    doc.save(out)
    print("written", out)


if __name__ == "__main__":
    main()
